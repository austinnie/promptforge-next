# skills/image_curator/writers.py
"""image_curator 多格式输出器。

    WordWriter      → .docx（python-docx）
    PrintWriter     → .print.html（浏览器 Ctrl+P 存 PDF）
    ClipboardWriter → 内联样式 HTML，一键复制到微信/知乎编辑器
"""

import html
import logging
import re
from datetime import datetime
from pathlib import Path
from typing import Any

logger = logging.getLogger(__name__)


def make_heading(entry: dict[str, Any]) -> str:
    """生成友好的文章小节标题。"""
    name = Path(entry["filename"]).stem

    m = re.match(r"^作品\s*(\d+)$", name)
    if m:
        return f"作品 {int(m.group(1)):02d}"

    # 规则 2：清理常见前缀
    name = re.sub(r"^\d{8}[_\-\s]*", "", name)       # 日期前缀
    name = re.sub(r"^(api|text2img|img2img)[_\-\s]*", "", name)

    # 规则 3：去掉结尾的 8 位 hex（新的命名规则：{摘要}_{8位hex}）
    name = re.sub(r"[_\-\s]+[0-9a-f]{8}$", "", name, flags=re.IGNORECASE)
    # 去掉纯 32 位 hex（旧的命名规则）
    name = re.sub(r"^[0-9a-f]{32}$", "", name, flags=re.IGNORECASE)

    name = name.replace("_", " ").strip()

    # 规则 4：兜底
    if not name or len(name) < 3 or name.isdigit():
        return f"作品 {entry['index']:02d}"
        
    return name[:40]


# ============================================================
# Word
# ============================================================

class WordWriter:
    """生成 .docx（需 python-docx）"""

    def __init__(self, font_name: str = "Microsoft YaHei",
                 east_asia_font: str = "微软雅黑"):
        self.font_name = font_name
        self.east_asia_font = east_asia_font

    def write(self, path, title: str, intro: str, entries: list[dict]) -> bool:
        try:
            from docx import Document
            from docx.shared import Pt, Cm, RGBColor
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            from docx.oxml.ns import qn
        except ImportError:
            logger.warning("⚠️ 未安装 python-docx，跳过 Word")
            return False

        doc = Document()

        normal = doc.styles["Normal"]
        normal.font.name = self.font_name
        normal.font.size = Pt(11)
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), self.east_asia_font)

        h = doc.add_heading(title, level=0)
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in h.runs:
            r.font.name = self.font_name
            r._element.rPr.rFonts.set(qn("w:eastAsia"), self.east_asia_font)

        meta = doc.add_paragraph()
        meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = meta.add_run(
            f"共 {len(entries)} 张作品 · {datetime.now().strftime('%Y-%m-%d')}"
        )
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

        intro_p = doc.add_paragraph(intro)
        intro_p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        doc.add_paragraph()

        for e in entries:
            heading = make_heading(e)
            h2 = doc.add_heading(f"{e['index']}. {heading}", level=2)
            for r in h2.runs:
                r.font.name = self.font_name
                r._element.rPr.rFonts.set(qn("w:eastAsia"), self.east_asia_font)

            img_path = Path(e.get("local_path") or e["source_path"])
            if img_path.exists():
                try:
                    doc.add_picture(str(img_path), width=Cm(14))
                    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                except Exception as ex:
                    logger.warning(f"Word 插图失败 {img_path.name}: {ex}")

            p = doc.add_paragraph(e["description"])
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

            fname = doc.add_paragraph()
            r = fname.add_run(e["filename"])
            r.font.size = Pt(8)
            r.font.color.rgb = RGBColor(0xAA, 0xAA, 0xAA)

            doc.add_paragraph()

        doc.save(str(path))
        return True


# ============================================================
# 打印版 HTML（用于 Ctrl+P 存 PDF）
# ============================================================

_PRINT_CSS = """
@page { size: A4; margin: 18mm 16mm; }
body { font-family: "Songti SC", "STSong", "SimSun", serif;
       color: #2a2a2a; line-height: 1.75; font-size: 11pt; }
h1 { font-size: 22pt; text-align: center; letter-spacing: 2pt; margin: 0 0 6pt; }
.meta { text-align: center; color: #999; font-size: 9pt; margin: 0 0 24pt; }
.intro { border-top: 1px solid #ddd; border-bottom: 1px solid #ddd;
         padding: 12pt 0; margin: 0 0 28pt; color: #555;
         font-style: italic; text-align: justify; }
.entry { margin-bottom: 32pt; page-break-inside: avoid; }
.entry h2 { font-size: 14pt; margin: 0 0 12pt; padding-left: 8pt;
            border-left: 2pt solid #c25b3b; page-break-after: avoid; }
.entry img { max-width: 100%; max-height: 420pt; display: block;
             margin: 0 auto 12pt; }
.entry .desc { text-align: justify; margin: 0 0 8pt; }
.entry .filename { color: #bbb; font-size: 8pt;
                   font-family: Consolas, monospace; word-break: break-all; }
"""


class PrintWriter:
    """生成浏览器打印版 HTML（Ctrl+P → 另存为 PDF）。"""

    def write(self, path, title: str, intro: str, entries: list[dict]) -> bool:
        rows = []
        for e in entries:
            heading = html.escape(make_heading(e))
            # 用相对 assets 路径（跟 article.html 在同一目录下）
            img_src = html.escape(e.get("asset_path") or "")
            desc = html.escape(e["description"])
            fname = html.escape(e["filename"])
            rows.append(
                f'<div class="entry">'
                f'<h2>{e["index"]}. {heading}</h2>'
                f'<img src="{img_src}" alt="{heading}">'
                f'<p class="desc">{desc}</p>'
                f'<p class="filename">{fname}</p>'
                f'</div>'
            )

        doc = (
            f'<!doctype html><html><head><meta charset="utf-8">'
            f'<title>{html.escape(title)}</title>'
            f'<style>{_PRINT_CSS}</style></head><body>'
            f'<h1>{html.escape(title)}</h1>'
            f'<p class="meta">共 {len(entries)} 张作品 · '
            f'{datetime.now().strftime("%Y-%m-%d")}</p>'
            f'<div class="intro">{html.escape(intro)}</div>'
            + "".join(rows) +
            '</body></html>'
        )
        Path(path).write_text(doc, encoding="utf-8")
        return True


# ============================================================
# 富文本（微信 / 知乎）
# ============================================================

class ClipboardWriter:
    """内联样式 HTML，用户点「复制全文」→ 粘贴到微信/知乎。"""

    def write(self, path, title: str, intro: str, entries: list[dict]) -> bool:
        content = self._build_inline(title, intro, entries)

        page = f"""<!doctype html>
<html lang="zh-CN">
<head>
<meta charset="utf-8">
<title>复制到微信/知乎 · {html.escape(title)}</title>
<style>
  body {{ margin:0; font-family:-apple-system,"Microsoft YaHei",sans-serif;
          background:#f5f5f5; color:#2a2a2a; }}
  .toolbar {{ position:sticky; top:0; background:#c25b3b; color:#fff;
              padding:12px 20px; display:flex; align-items:center; gap:16px;
              box-shadow:0 2px 8px rgba(0,0,0,.15); z-index:10; }}
  .toolbar button {{ padding:8px 20px; border:0; border-radius:4px; cursor:pointer;
                     background:#fff; color:#c25b3b; font-weight:bold; font-size:14px; }}
  .toolbar button:hover {{ background:#ffe8e0; }}
  .toolbar .hint {{ font-size:13px; opacity:.92; }}
  .container {{ max-width:680px; margin:20px auto 60px; background:#fff;
                padding:36px 30px; border-radius:6px;
                box-shadow:0 2px 16px rgba(0,0,0,.06); }}
</style>
</head>
<body>
<div class="toolbar">
  <button onclick="copyAll()">📋 复制全文</button>
  <span class="hint">粘贴到微信/知乎编辑器（图片需手动上传）</span>
</div>
<div class="container" id="content">
{content}
</div>
<script>
async function copyAll() {{
  const el = document.getElementById('content');
  const htmlStr = el.innerHTML;
  const textStr = el.innerText;
  try {{
    const item = new ClipboardItem({{
      'text/html':  new Blob([htmlStr], {{type:'text/html'}}),
      'text/plain': new Blob([textStr], {{type:'text/plain'}})
    }});
    await navigator.clipboard.write([item]);
    alert('✅ 已复制！请到微信/知乎编辑器粘贴。');
  }} catch(e) {{
    const range = document.createRange();
    range.selectNodeContents(el);
    const sel = window.getSelection();
    sel.removeAllRanges(); sel.addRange(range);
    document.execCommand('copy');
    alert('✅ 已复制（兼容模式）。');
  }}
}}
</script>
</body>
</html>"""
        Path(path).write_text(page, encoding="utf-8")
        return True

    def _build_inline(self, title: str, intro: str, entries: list[dict]) -> str:
        parts = []
        parts.append(
            f'<h1 style="text-align:center;font-size:24px;color:#2a2a2a;'
            f'letter-spacing:2px;margin:0 0 12px;font-weight:700;">'
            f'{html.escape(title)}</h1>'
        )
        parts.append(
            f'<p style="text-align:center;color:#999;font-size:13px;margin:0 0 24px;">'
            f'共 {len(entries)} 张作品 · {datetime.now().strftime("%Y-%m-%d")}</p>'
        )
        parts.append(
            f'<section style="border-top:1px solid #eee;border-bottom:1px solid #eee;'
            f'padding:20px 8px;margin:24px 0;color:#555;font-style:italic;'
            f'font-size:15px;line-height:1.85;text-align:justify;">'
            f'{html.escape(intro)}</section>'
        )

        for e in entries:
            heading = html.escape(make_heading(e))
            asset = e.get("asset_path") or ""
            parts.append(
                f'<h2 style="font-size:18px;color:#2a2a2a;margin:40px 0 16px;'
                f'padding-left:12px;border-left:3px solid #c25b3b;">'
                f'{e["index"]}. {heading}</h2>'
            )
            if asset:
                parts.append(
                    f'<p style="text-align:center;margin:0 0 20px;">'
                    f'<img src="{html.escape(asset)}" '
                    f'style="max-width:100%;border-radius:4px;display:block;margin:0 auto;">'
                    f'</p>'
                )
            parts.append(
                f'<p style="font-size:15px;line-height:1.9;color:#333;'
                f'text-align:justify;margin:0 0 12px;">'
                f'{html.escape(e["description"])}</p>'
            )
            parts.append(
                f'<p style="font-size:12px;color:#bbb;font-family:monospace;'
                f'word-break:break-all;margin:0 0 32px;">'
                f'{html.escape(e["filename"])}</p>'
            )

        return "\n".join(parts)