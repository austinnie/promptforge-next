# skills/tech_hot_article/skill.py
"""技术热点文章生成 skill。

能力：
    generate    抓取热点 → LLM 写稿 → AI 配图 → 导出 Word
    list_hot    只看热点列表

改造点：
    - Ollama → dispatcher.chat
    - api_engines → dispatcher.generate_image
    - 输出目录 → data/assets/tech_article/
    - 接口全部 async
"""

import asyncio
import json
import logging
import random
import re
import time
from datetime import datetime
from pathlib import Path
from typing import Any, Optional

logger = logging.getLogger(__name__)

try:
    import feedparser
    FEEDPARSER_AVAILABLE = True
except ImportError:
    FEEDPARSER_AVAILABLE = False

try:
    from docx import Document
    from docx.shared import Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

try:
    from PIL import Image, ImageDraw, ImageFont
    PIL_AVAILABLE = True
except ImportError:
    PIL_AVAILABLE = False


# ==================== 常量 ====================

HOT_SOURCES = {
    "hacker_news": "https://hnrss.org/frontpage",
    "techcrunch": "https://techcrunch.com/feed/",
    "the_verge": "https://www.theverge.com/rss/index.xml",
    "wired": "https://www.wired.com/feed/rss",
    "arstechnica": "https://arstechnica.com/feed/",
    "devto": "https://dev.to/feed",
}

WRITING_STYLES = [
    "专业分析型",
    "通俗科普型",
    "深度技术型",
    "行业观察型",
    "趋势预测型",
]


class TechHotArticle:
    """技术热点文章生成器。"""

    NAME = "tech_hot_article"
    VERSION = "2.0.0"

    def __init__(self, config: Optional[dict] = None):
        self.config = config or {}
        base = Path(self.config.get("output_dir", "./data/assets/tech_article"))
        self.output_dir = base
        self.images_dir = base / "images"
        self.word_dir = base / "word"
        self.meta_dir = base / "meta"

        for d in (self.output_dir, self.images_dir, self.word_dir, self.meta_dir):
            d.mkdir(parents=True, exist_ok=True)

        self.max_hot_items = int(self.config.get("max_hot_items", 10))
        self.article_words = int(self.config.get("article_words", 1500))
        self.temperature = float(self.config.get("temperature", 0.85))
        self.image_width = int(self.config.get("image_width", 800))
        self.image_height = int(self.config.get("image_height", 600))
        self.chars_per_image = int(self.config.get("chars_per_image", 500))
        self.max_images = int(self.config.get("max_images_per_article", 5))
        self.min_images = int(self.config.get("min_images_per_article", 1))
        self.image_style = self.config.get(
            "image_style",
            "digital art, tech illustration, clean, modern, professional",
        )

    # ==================== 主入口 ====================

    async def execute(self, **kwargs) -> dict:
        action = kwargs.pop("action", "generate")
        try:
            if action == "list_hot":
                return self._ok(await self._action_list_hot(**kwargs))
            if action == "generate":
                return self._ok(await self._action_generate(**kwargs))
            return self._err(f"未知 action: {action}")
        except Exception as e:
            logger.exception("TechHotArticle 执行失败")
            return self._err(str(e))

    async def _action_list_hot(self, **kwargs) -> dict:
        if not FEEDPARSER_AVAILABLE:
            raise RuntimeError("feedparser 未安装，请 pip install feedparser")
        items = self.get_hot_topics()
        return {"hot_items": items, "count": len(items)}

    async def _action_generate(
        self,
        hot_index: Optional[int] = None,
        style: str = "",
        article_words: Optional[int] = None,
        with_images: bool = True,
        language: str = "zh",
    ) -> dict:
        start = time.time()

        if not FEEDPARSER_AVAILABLE:
            raise RuntimeError("feedparser 未安装")
        if not DOCX_AVAILABLE:
            raise RuntimeError("python-docx 未安装")

        # 1. 抓热点
        logger.info("获取技术热点...")
        hot_items = self.get_hot_topics()
        if not hot_items:
            raise RuntimeError("未获取到任何热点")

        if hot_index is not None and 0 <= hot_index < len(hot_items):
            hot = hot_items[hot_index]
        else:
            hot = random.choice(hot_items)

        # 2. 选风格
        if style not in WRITING_STYLES:
            style = random.choice(WRITING_STYLES)

        # 3. 写文章
        logger.info(f"生成文章，风格={style}")
        article = await self._generate_article(
            hot, style, article_words or self.article_words,
        )
        if not article:
            raise RuntimeError("文章生成失败")

        # 4. 拆分段落 + 规划配图位置
        paragraphs = self._split_paragraphs(article["body"])
        image_positions_idx = self._determine_image_positions(paragraphs)

        # 5. 逐位置生成配图
        image_positions: dict[int, str] = {}
        if with_images:
            for idx in image_positions_idx:
                try:
                    img_path = await self._generate_image_for_block(
                        paragraphs[idx], article["title"],
                    )
                    if img_path:
                        image_positions[idx] = img_path
                except Exception as e:
                    logger.warning(f"配图生成失败（位置 {idx}）: {e}")

        # 6. Word 导出
        word_path = await asyncio.to_thread(
            self._create_word_document, article, image_positions,
        )

        # 7. 元数据 JSON
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        meta_file = self.meta_dir / f"article_{ts}.json"
        meta_data = {
            "article": article,
            "hot_item": hot,
            "image_paths": list(image_positions.values()),
            "image_positions": {str(k): v for k, v in image_positions.items()},
            "word_path": word_path,
            "timestamp": datetime.now().isoformat(),
        }
        meta_file.write_text(
            json.dumps(meta_data, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )

        result = {
            "title": article["title"],
            "hot_topic": hot["title"],
            "hot_source": hot.get("source", ""),
            "style": style,
            "body": article["body"],
            "word_file": self._to_url(word_path),
            "image_files": [self._to_url(p) for p in image_positions.values()],
            "article_file": self._to_url(str(meta_file)),
            "image_count": len(image_positions),
            "paragraph_count": len(paragraphs),
            "word_count": len(article["body"]),
            "generated_at": datetime.now().isoformat(),
            "elapsed": f"{time.time() - start:.1f}s",
        }
        logger.info(f"✅ 完成: {result['title']} ({result['image_count']} 图)")
        return result

    # ==================== LLM ====================

    async def _call_llm(
        self, prompt: str, temperature: float = 0.85,
        max_tokens: int = 4096, system: str = "",
    ) -> str:
        from engines import dispatcher

        messages = []
        if system:
            messages.append({"role": "system", "content": system})
        messages.append({"role": "user", "content": prompt})

        try:
            text, used = await dispatcher.chat(
                messages=messages,
                temperature=temperature,
                max_tokens=max_tokens,
            )
            logger.info(f"[llm] 引擎={used} 返回 {len(text)} 字")
            return text.strip()
        except Exception as e:
            logger.error(f"[llm] 失败: {e}")
            return ""

    async def _generate_article(
        self, hot: dict, style: str, words: int,
    ) -> Optional[dict]:
        title = hot.get("title", "")
        summary = hot.get("summary", "")
        source = hot.get("source", "")

        prompt = f"""你是一位资深科技编辑，请根据以下热点信息，撰写一篇技术文章。

热点标题：{title}
热点描述：{summary}
来源：{source}

【重要】无论热点标题是中文还是英文，文章标题和正文都必须使用中文。热点标题仅作为素材，不要直接照搬英文原标题。

写作风格：{style}
文章字数：约 {words} 字

要求：
1. 第一行是**中文**文章标题（不要加"标题："前缀，不要加 # 号）
2. 空一行后开始正文
3. 正文用中文，分段落，每段之间用空行分隔
4. 分析技术背景、影响和未来趋势
5. 语言专业但不晦涩
6. 不要使用 Markdown 标记（#、**、---等）

请直接输出文章："""

        content = await self._call_llm(prompt, self.temperature, max_tokens=4096)
        if not content:
            return None

        content = self._clean_markdown(content)
        lines = content.split("\n")
        art_title = lines[0].strip().lstrip("#").strip()
        if not art_title:
            art_title = title

        body = "\n".join(lines[1:]).strip()

        return {
            "title": art_title,
            "body": body,
            "full_content": content,
            "style": style,
            "hot_title": title,
            "hot_source": source,
        }

    # ==================== 热点抓取 ====================

    def get_hot_topics(self) -> list[dict]:
        all_items: list[dict] = []
        all_items += self._fetch_rss("hacker_news", "Hacker News")
        all_items += self._fetch_rss("devto", "Dev.to")
        all_items += self._fetch_rss("techcrunch", "TechCrunch")
        all_items += self._fetch_rss("the_verge", "The Verge")

        if not all_items:
            logger.warning("所有源抓取失败，使用模拟数据")
            all_items = self._mock_hot_topics()

        # 去重
        seen: set[str] = set()
        unique: list[dict] = []
        for it in all_items:
            key = it["title"][:30].lower()
            if key not in seen:
                seen.add(key)
                unique.append(it)

        unique.sort(key=lambda x: x.get("score", 0), reverse=True)
        return unique[: self.max_hot_items]

    def _fetch_rss(self, key: str, source_name: str) -> list[dict]:
        url = HOT_SOURCES.get(key)
        if not url:
            return []
        try:
            feed = feedparser.parse(url)
            items = []
            for entry in feed.entries[:10]:
                items.append({
                    "title": entry.get("title", ""),
                    "link": entry.get("link", ""),
                    "summary": self._clean_html(entry.get("summary", ""))[:300],
                    "source": source_name,
                    "score": 0,
                })
            logger.info(f"✅ {source_name}: {len(items)} 条")
            return items
        except Exception as e:
            logger.warning(f"抓取 {source_name} 失败: {e}")
            return []

    @staticmethod
    def _clean_html(text: str) -> str:
        text = re.sub(r"<[^>]+>", " ", text or "")
        text = re.sub(r"&\w+;", " ", text)
        text = re.sub(r"\s+", " ", text)
        return text.strip()

    @staticmethod
    def _mock_hot_topics() -> list[dict]:
        return [
            {"title": "ChatGPT 发布多模态功能",
             "link": "", "summary": "OpenAI 新增图像识别能力",
             "source": "TechCrunch", "score": 950},
            {"title": "Google 发布 Gemini 2.0",
             "link": "", "summary": "新一代 AI 模型性能领先",
             "source": "Wired", "score": 890},
            {"title": "Meta 开源 Llama 3",
             "link": "", "summary": "70B 模型免费商用",
             "source": "The Verge", "score": 850},
            {"title": "苹果发布 Apple Intelligence",
             "link": "", "summary": "Siri 获 AI 能力提升",
             "source": "ZDNet", "score": 800},
            {"title": "量子计算重大突破",
             "link": "", "summary": "Google 实现量子霸权 2.0",
             "source": "Ars Technica", "score": 750},
        ]

    # ==================== 段落 / 配图位置 ====================

    @staticmethod
    def _split_paragraphs(body: str) -> list[str]:
        parts = [p.strip() for p in body.split("\n\n") if p.strip()]
        if len(parts) < 3:
            parts = [p.strip() for p in body.split("\n")
                     if p.strip() and len(p.strip()) > 20]
        return [p for p in parts if len(p) > 30]

    def _determine_image_positions(self, paragraphs: list[str]) -> list[int]:
        total_p = len(paragraphs)
        if total_p == 0:
            return []
        total_chars = sum(len(p) for p in paragraphs)

        calculated = max(1, total_chars // self.chars_per_image)
        count = max(self.min_images, min(calculated, self.max_images))
        count = min(count, max(0, total_p - 1))
        if count <= 0:
            return []

        step = total_p / (count + 1)
        positions = []
        for i in range(1, count + 1):
            pos = int(step * i)
            if pos < total_p:
                positions.append(pos)
        logger.info(f"📊 {total_p} 段 {total_chars} 字 → {len(positions)} 图")
        return positions

    # ==================== 图片生成 ====================

    async def _generate_image_for_block(
        self, block_text: str, article_title: str,
    ) -> Optional[str]:
        from engines import dispatcher

        snippet = block_text[:80].replace("\n", " ").strip()
        prompt = (
            f"Tech article illustration about: {snippet}. "
            f"Context: {article_title[:50]}. "
            f"Style: {self.image_style}, vibrant colors, professional, "
            f"no text, no watermark"
        )

        try:
            img, used = await dispatcher.generate_image(
                prompt=prompt,
                negative="text, watermark, signature, low quality, blurry",
                width=self.image_width,
                height=self.image_height,
                seed=random.randint(1, 999),
            )
            logger.info(f"🎨 配图引擎={used}")

            ts = datetime.now().strftime("%Y%m%d_%H%M%S_%f")[:19]
            path = self.images_dir / f"img_{ts}.png"

            target = (self.image_width, self.image_height)
            if img.size != target:
                img = img.resize(target, Image.Resampling.LANCZOS)
            img.save(path, "PNG")
            return str(path)

        except Exception as e:
            logger.warning(f"配图失败，回退 Pillow: {e}")
            return self._pillow_fallback(article_title)

    def _pillow_fallback(self, title: str) -> Optional[str]:
        if not PIL_AVAILABLE:
            return None

        ts = datetime.now().strftime("%Y%m%d_%H%M%S_%f")[:19]
        path = self.images_dir / f"fallback_{ts}.png"

        img = Image.new("RGB", (self.image_width, self.image_height), (20, 30, 50))
        draw = ImageDraw.Draw(img)

        colors = [(50, 100, 200), (200, 50, 100), (50, 200, 100),
                  (200, 150, 50), (150, 50, 200)]
        for _ in range(random.randint(3, 6)):
            x1 = random.randint(0, self.image_width)
            y1 = random.randint(0, self.image_height)
            draw.rectangle(
                [x1, y1, x1 + 200, y1 + 200],
                outline=random.choice(colors), width=3,
            )

        try:
            font = None
            for fp in ["C:/Windows/Fonts/msyh.ttc",
                       "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"]:
                if Path(fp).exists():
                    font = ImageFont.truetype(fp, 24)
                    break
            if font is None:
                font = ImageFont.load_default()
            text = title[:40]
            bbox = draw.textbbox((0, 0), text, font=font)
            draw.text(
                ((self.image_width - (bbox[2] - bbox[0])) // 2,
                 (self.image_height - (bbox[3] - bbox[1])) // 2),
                text, fill=(255, 255, 255), font=font,
            )
        except Exception:
            pass

        img.save(path, "PNG")
        return str(path)

    # ==================== Word ====================

    def _create_word_document(
        self, article: dict, image_positions: dict[int, str],
    ) -> str:
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        safe = re.sub(r'[<>:"/\\|?*]', "", article["title"])[:30]
        path = self.word_dir / f"{safe}_{ts}.docx"

        doc = Document()
        title = doc.add_heading(article["title"], 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        meta = doc.add_paragraph()
        meta.add_run(f"生成时间：{datetime.now():%Y-%m-%d %H:%M:%S}\n")
        meta.add_run(f"写作风格：{article.get('style', '')}\n")
        meta.add_run(f"热点来源：{article.get('hot_source', '')}")
        doc.add_paragraph()

        paragraphs = self._split_paragraphs(article.get("body", ""))
        for idx, para in enumerate(paragraphs):
            p = doc.add_paragraph(para)
            p.paragraph_format.first_line_indent = Inches(0.3)

            if idx in image_positions:
                img_path = image_positions[idx]
                if img_path and Path(img_path).exists():
                    try:
                        doc.add_picture(img_path, width=Inches(5.5))
                        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                    except Exception as e:
                        logger.warning(f"插图失败: {e}")

        doc.add_page_break()
        footer = doc.add_paragraph()
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.add_run(f"由 PromptForge 自动生成 | {datetime.now():%Y-%m-%d}")

        doc.save(str(path))
        logger.info(f"✅ Word: {path}")
        return str(path)

    # ==================== 工具 ====================

    @staticmethod
    def _clean_markdown(text: str) -> str:
        text = re.sub(r"```.*?```", "", text, flags=re.DOTALL)
        text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
        text = re.sub(r"\*(.+?)\*", r"\1", text)
        text = re.sub(r"^#{1,6}\s+", "", text, flags=re.MULTILINE)
        text = re.sub(r"^[-=*]{3,}\s*$", "", text, flags=re.MULTILINE)
        return text.strip()

    def _to_url(self, local_path: str) -> str:
        """data/assets/xxx → /files/xxx"""
        p = Path(local_path).resolve()
        root = Path("./data/assets").resolve()
        try:
            rel = p.relative_to(root).as_posix()
            return f"/files/{rel}"
        except ValueError:
            return local_path

    def _ok(self, result: dict) -> dict:
        return {
            "status": "success",
            "result": result,
            "skill": self.NAME,
            "version": self.VERSION,
            "executed_at": datetime.now().isoformat(timespec="seconds"),
        }

    def _err(self, msg: str) -> dict:
        return {
            "status": "error",
            "error": msg,
            "skill": self.NAME,
            "version": self.VERSION,
            "executed_at": datetime.now().isoformat(timespec="seconds"),
        }